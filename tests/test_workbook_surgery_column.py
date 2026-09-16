"""Opt-in workbook integration: CODEX_NODE_MODULES=... NODE_BINARY=... python -m unittest ..."""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl
from test_disease_specific_plans import disease_plan, extracted, group, medication, patient, profile

SKILL = Path(__file__).resolve().parents[1]


@unittest.skipUnless(os.environ.get("CODEX_NODE_MODULES"), "需要artifact-tool运行环境")
class WorkbookSurgeryColumnTest(unittest.TestCase):
    def run_script(self, script, *args, success=True):
        binary = os.environ.get("NODE_BINARY", "node") if script.endswith(".mjs") else sys.executable
        result = subprocess.run([binary, str(SKILL / "scripts" / script), *map(str, args)],
                                capture_output=True, text=True, env={**os.environ, "AUTO_MEDICATION_SEARCH": "0"})
        if success:
            self.assertEqual(result.returncode, 0, result.stderr[-4000:])
        return result

    def test_device_and_drug_saved_columns_roundtrip_and_tamper_rejection(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for product_type in ["器械", "用药"]:
                with self.subTest(product_type=product_type):
                    folder = root / product_type
                    folder.mkdir()
                    people = [patient("test-a", "测试疾病甲"), patient("test-b", "测试疾病乙")]
                    for p in people:
                        p["activateTime"] = ""
                        p["sourceConfirmationTime"] = "2026-04-10 12:00:00"
                    people[1]["adverseEvent"] = "是"
                    data = extracted(people)
                    data["inputFormat"] = "medicationReminder14"
                    plans = [disease_plan(d, [d], [group("组", [medication("测试候选药")])],
                                          minimum_combined=1, minimum_disease=0,
                                          count_rationale="仅测试格式映射，非临床方案")
                             for d in ["测试疾病甲", "测试疾病乙"]]
                    if product_type == '器械':
                        for plan in plans:
                            plan['minimumCombinedMedicationCount'] = 3
                            plan['minimumDiseaseMedicationCount'] = 3
                            plan['medicationGroups'] = [group(f'组{i}', [medication(f'测试药{i}')]) for i in range(3)]
                    spec = profile(plans, productType=product_type,
                                   baseMedication=medication("测试产品") if product_type == "用药" else None,
                                   surgeryRules=[],
                                   simulatedSurgeryRules=[{"when": {"diseaseEqualsAny": [d]}, "surgeryName": s,
                                                          "assumptions":["测试假设，非事实"],"rationale":"缺少实际术式资料，使用情景规则"}
                                                         for d, s in [("测试疾病甲", "测试甲术式"), ("测试疾病乙", "测试乙术式")]])
                    for name, value in [("patients", data), ("profile", spec)]:
                        (folder / f"{name}.json").write_text(json.dumps(value, ensure_ascii=False))
                    payload_path = folder / "payload.json"
                    self.run_script("generate_payload.py", "--patients", folder / "patients.json", "--profile", folder / "profile.json", "--output", payload_path)
                    self.run_script("build_workbooks.mjs", "--payload", payload_path,
                                    "--reminder-template", SKILL / "assets/medication-reminder-template.xlsx",
                                    "--medication-template", SKILL / "assets/medication-list-template.xlsx", "--output-dir", folder)
                    suffix = "_模拟" if product_type == "器械" else ""
                    reminder = folder / f"用药提醒_测试产品{suffix}.xlsx"
                    meds = folder / f"用药方案_测试产品{suffix}.xlsx"
                    generated = json.loads(payload_path.read_text())
                    self.assertEqual(generated['meta']['simulation'], product_type == '器械')
                    self.assertEqual(len(generated['meta']['surgerySimulationAudit']), 2 if product_type == '器械' else 0)
                    verify_args = ["--payload", payload_path, "--reminder", reminder, "--medication", meds, "--report", folder / "verify.json"]
                    self.run_script("verify_workbooks.mjs", *verify_args)
                    workbook = openpyxl.load_workbook(reminder, data_only=True)
                    sheet = workbook.active
                    self.assertEqual(sheet.max_column, 14)
                    self.assertFalse(any('模拟' in str(cell.value or '') for row in sheet for cell in row))
                    self.assertEqual([sheet.cell(2, c).value for c in (7, 8, 9)], ["既往过敏史", "手术名称", "联合用药"])
                    self.assertEqual([sheet.cell(r, 8).value or "" for r in (3, 4)],
                                     ["测试甲术式", "测试乙术式"] if product_type == "器械" else ["", ""])
                    self.assertEqual(sheet["N4"].value, "是")
                    self.assertEqual(sheet["J3"].value, "2026-04-10 12:00:00")
                    self.assertIn("A1:N1", {str(r) for r in sheet.merged_cells.ranges})
                    self.assertEqual(next(iter(sheet.tables.values())).ref, "A2:N4")
                    workbook.close()
                    self.run_script("extract_patients.py", "--source", reminder, "--output", folder / "roundtrip.json")
                    roundtrip = json.loads((folder / "roundtrip.json").read_text())
                    self.assertEqual(roundtrip["inputFormat"], "medicationReminder14")
                    self.assertEqual([p["userid"] for p in roundtrip["patients"]], ["test-a", "test-b"])
                    self.assertEqual(roundtrip["patients"][0]["sourceConfirmationTime"], "2026-04-10 12:00:00")
                    if os.environ.get("WORKBOOK_TEST_ARTIFACTS"):
                        destination = Path(os.environ["WORKBOOK_TEST_ARTIFACTS"]) / product_type
                        destination.mkdir(parents=True, exist_ok=True)
                        for artifact in [reminder, meds, *folder.glob("*.png")]:
                            shutil.copy2(artifact, destination / artifact.name)
                    # Corrupt only the saved surgery cell; payload remains valid.
                    corrupted = folder / "corrupted.xlsx"
                    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                    with zipfile.ZipFile(reminder) as source, zipfile.ZipFile(corrupted, "w") as dest:
                        for info in source.infolist():
                            content = source.read(info.filename)
                            if info.filename == "xl/worksheets/sheet1.xml":
                                xml = ET.fromstring(content)
                                row = xml.find(f"{ns}sheetData/{ns}row[@r='3']")
                                cell = row.find(f"{ns}c[@r='H3']")
                                if cell is None:
                                    cell = ET.SubElement(row, f"{ns}c", {"r": "H3"})
                                for child in list(cell):
                                    cell.remove(child)
                                cell.set("t", "inlineStr")
                                ET.SubElement(ET.SubElement(cell, f"{ns}is"), f"{ns}t").text = "错误术式"
                                content = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
                            dest.writestr(info, content)
                    verify_args[3] = corrupted
                    result = self.run_script("verify_workbooks.mjs", *verify_args, success=False)
                    self.assertNotEqual(result.returncode, 0)
                    self.assertIn("手术名称与生成记录不一致", result.stderr)
