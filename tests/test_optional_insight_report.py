import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import json
import re
import tempfile
import unittest
from pathlib import Path

from docx import Document

from build_insight_report import build_report
from extract_insight_sources import build_insight
from generate_insight_charts import generate_charts
import test_extract_insight_sources as source_tests
from validate_insight_report import validate_report


class OptionalAdverseReportTest(unittest.TestCase):
    def setUp(self):
        self.fixture = source_tests.ExtractInsightSourcesTest()
        self.fixture.setUp()
        self.addCleanup(self.fixture.tearDown)
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def build_report_for(self, paths):
        insight = build_insight(paths, "注射用胰蛋白酶", "2026-07-01", "2026-07-31")
        insight_path = self.root / "insight.json"
        insight_path.write_text(json.dumps(insight, ensure_ascii=False))
        chart_path = generate_charts(insight, self.root / "charts")
        output = self.root / "report.docx"
        build_report(insight_path, chart_path, None, output)
        self.assertEqual(validate_report(output, insight_path), [])
        layout = Document(output)
        for paragraph in layout.paragraphs:
            if paragraph._p.xpath('.//w:drawing'):
                self.assertTrue(paragraph.paragraph_format.keep_with_next)
        for table in layout.tables[1:]:
            for cell in table.rows[-1].cells:
                self.assertTrue(cell.paragraphs[-1].paragraph_format.keep_with_next)
        return insight, json.loads(Path(chart_path).read_text()), Document(output), output, insight_path

    def test_six_file_pipeline_omits_event_and_risk_statistics(self):
        paths = [path for path in self.fixture.paths if not path.endswith("ae.xlsx")]
        insight, charts, document, output, insight_path = self.build_report_for(paths)
        self.assertFalse(insight['metrics']['adverseEvents']['provided'])
        self.assertNotIn('adverse-events', {chart['id'] for chart in charts})
        self.assertNotIn('risk', {chart['id'] for chart in charts})
        self.assertEqual(len([p for p in document.paragraphs if re.match(r'^[一二三四五六七八九]、', p.text)]), 9)
        labels = [row.cells[0].text for table in document.tables for row in table.rows]
        self.assertNotIn('不良反应发生率', labels)
        self.assertNotIn('不良反应监测', labels)
        self.assertNotIn('低风险', labels)
        text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
        self.assertNotRegex(text, r'不良反应0例|低风险1人|未记录不良反应的患者为1人')
        # Validation must catch a false zero-event assertion in optional mode.
        document.add_paragraph('本周期记录不良反应0例，患者发生率为0.0%。')
        document.save(output)
        self.assertIn('不良反应清单选填模式不得输出事件统计或风险分层结论', validate_report(output, insight_path))

    def test_provided_empty_workbook_preserves_zero_event_statistics(self):
        ae_path = next(Path(path) for path in self.fixture.paths if path.endswith('ae.xlsx'))
        source_tests.write_book(ae_path, ['患者ID', '不良反应发生时间', '不良反应严重程度分级'], [])
        insight, charts, document, _, _ = self.build_report_for(self.fixture.paths)
        self.assertTrue(insight['metrics']['adverseEvents']['provided'])
        self.assertEqual(insight['metrics']['adverseEventRate']['numerator'], 0)
        self.assertEqual(insight['metrics']['adverseEventRate']['denominator'], 1)
        self.assertIn('adverse-events', {chart['id'] for chart in charts})
        self.assertIn('risk', {chart['id'] for chart in charts})
        text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn('本周期记录不良反应0例', text)

    def test_seven_file_pipeline_retains_recorded_events(self):
        insight, charts, document, _, _ = self.build_report_for(self.fixture.paths)
        self.assertEqual(insight['metrics']['adverseEvents']['recordCount'], 1)
        self.assertEqual(insight['metrics']['adverseEventRate']['value'], 1)
        self.assertIn('adverse-events', {chart['id'] for chart in charts})
        self.assertIn('本周期记录不良反应1例', '\n'.join(paragraph.text for paragraph in document.paragraphs))


if __name__ == '__main__':
    unittest.main()
