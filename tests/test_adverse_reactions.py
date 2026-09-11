import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


SKILL_DIR = Path(__file__).resolve().parents[1]
SCRIPT = SKILL_DIR / "scripts" / "generate_adverse_reactions.py"
BUILDER = SKILL_DIR / "scripts" / "build_adverse_reaction_workbook.mjs"
VERIFIER = SKILL_DIR / "scripts" / "verify_adverse_reaction_workbook.mjs"
TEMPLATE = SKILL_DIR / "assets" / "adverse-reaction-template.xlsx"
NODE = os.environ.get("CODEX_NODE") or shutil.which("node")
NODE_MODULES = os.environ.get("CODEX_NODE_MODULES")


def patient(
    userid,
    tag,
    *,
    disease="脑梗死",
    activated="2026-04-10 10:00:00",
    gender="男",
    age=60,
    allergy="无",
):
    return {
        "sequence": 1,
        "userid": userid,
        "patientName": f"患者{userid}",
        "activateTime": activated,
        "gender": gender,
        "age": age,
        "disease": disease,
        "allergyHistory": allergy,
        "adverseEvent": "否",
        "adverseEventGrade": "",
        "patientTags": tag,
    }


def extracted(patients):
    return {
        "source": "/tmp/月度患者清单.xlsx",
        "title": "2026-04-月度患者清单",
        "headers": [],
        "patients": patients,
        "summary": {
            "patientCount": len(patients),
            "distinctUseridCount": len({item["userid"] for item in patients}),
            "activationMonths": {"2026-04": len(patients)},
        },
    }


class AdverseReactionGeneratorTest(unittest.TestCase):
    def run_generator(
        self, patients, product="血栓通胶囊", include_product=True,
        service_start="2026-04-01", service_end="2026-04-30",
    ):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            patients_path = temp / "patients.json"
            output_path = temp / "adverse-reactions.json"
            patients_path.write_text(json.dumps(extracted(patients), ensure_ascii=False), encoding="utf-8")
            command = [
                sys.executable,
                str(SCRIPT),
                "--patients",
                str(patients_path),
                "--output",
                str(output_path),
            ]
            if include_product:
                command.extend(["--product", product])
            if service_start is not None:
                command.extend(["--service-start", service_start])
            if service_end is not None:
                command.extend(["--service-end", service_end])
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                check=False,
            )
            payload = json.loads(output_path.read_text(encoding="utf-8")) if output_path.exists() else None
            return result, payload

    def test_filters_target_tags_and_preserves_userid_order(self):
        result, payload = self.run_generator(
            [
                patient("u0", "无"),
                patient("u1", "轻度患者"),
                patient("u2", "中度患者"),
                patient("u3", "重度患者"),
            ]
        )

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual([item["userid"] for item in payload["records"]], ["u1", "u2", "u3"])
        self.assertEqual([item["userid"] for item in payload["sourcePatients"]], ["u1", "u2", "u3"])
        self.assertEqual(payload["meta"]["sourcePatientCount"], 4)
        self.assertEqual(payload["meta"]["targetPatientCount"], 3)
        self.assertEqual(payload["meta"]["productName"], "血栓通胶囊")

    def test_requires_nonempty_product_name(self):
        missing_result, missing_payload = self.run_generator(
            [patient("u1", "中度患者")], include_product=False
        )
        blank_result, blank_payload = self.run_generator(
            [patient("u1", "中度患者")], product="   "
        )

        self.assertNotEqual(missing_result.returncode, 0)
        self.assertIsNone(missing_payload)
        self.assertIn("--product", missing_result.stderr)
        self.assertNotEqual(blank_result.returncode, 0)
        self.assertIsNone(blank_payload)
        self.assertIn("产品名称不能为空", blank_result.stderr)

    def test_maps_grade_and_manual_intervention_from_patient_tag(self):
        result, payload = self.run_generator(
            [
                patient("u1", "轻度患者"),
                patient("u2", "中度患者"),
                patient("u3", "重度患者"),
            ]
        )

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(payload["records"][0]["severityGrade"], "轻度")
        self.assertEqual(payload["records"][0]["manualIntervention"], "否")
        self.assertEqual(payload["records"][1]["severityGrade"], "中度")
        self.assertEqual(payload["records"][1]["manualIntervention"], "否")
        self.assertEqual(payload["records"][2]["severityGrade"], "重度")
        self.assertEqual(payload["records"][2]["manualIntervention"], "是")

    def test_time_follows_activation_within_service_period_and_is_deterministic(self):
        patients = [patient("stable-user", "中度患者", activated="2026-04-01 01:00:00")]

        first_result, first = self.run_generator(patients)
        second_result, second = self.run_generator(patients)

        self.assertEqual(first_result.returncode, 0, first_result.stderr)
        self.assertEqual(second_result.returncode, 0, second_result.stderr)
        self.assertEqual(first["records"], second["records"])
        occurrence = datetime.fromisoformat(first["records"][0]["occurrenceTime"])
        activation = datetime.fromisoformat(patients[0]["activateTime"])
        self.assertGreater(occurrence, activation)
        self.assertGreaterEqual(occurrence, datetime(2026, 4, 1))
        self.assertLessEqual(occurrence, datetime(2026, 4, 30, 23, 59, 59))
        self.assertGreaterEqual(occurrence.time(), datetime.strptime("07:30:00", "%H:%M:%S").time())
        self.assertLessEqual(occurrence.time(), datetime.strptime("21:59:59", "%H:%M:%S").time())
        self.assertEqual(first["meta"]["servicePeriod"], {"start": "2026-04-01", "end": "2026-04-30"})

    def test_occurrence_time_uses_daily_0730_to_2159_window(self):
        patients = [patient(f"daily-window-{index}", "中度患者", activated="2026-04-01 00:00:00") for index in range(100)]
        result, payload = self.run_generator(patients)
        self.assertEqual(result.returncode, 0, result.stderr)
        for record in payload["records"]:
            occurrence = datetime.fromisoformat(record["occurrenceTime"])
            self.assertGreaterEqual(occurrence.time(), datetime.strptime("07:30:00", "%H:%M:%S").time())
            self.assertLessEqual(occurrence.time(), datetime.strptime("21:59:59", "%H:%M:%S").time())

    def test_requires_both_service_dates(self):
        for options, flag in [
            ({"service_start": None}, "--service-start"),
            ({"service_end": None}, "--service-end"),
        ]:
            with self.subTest(flag=flag):
                result, payload = self.run_generator([patient("u1", "中度患者")], **options)
                self.assertNotEqual(result.returncode, 0)
                self.assertIsNone(payload)
                self.assertIn(flag, result.stderr)

    def test_rejects_invalid_or_reversed_service_period(self):
        for start, end in [
            ("", "2026-04-30"), ("2026-02-30", "2026-04-30"),
            ("2026-4-01", "2026-04-30"), ("2026-04-01", "2026-04-31"),
            ("2026-04-01T00:00:00", "2026-04-30"), ("2026-05-01", "2026-04-30"),
        ]:
            with self.subTest(start=start, end=end):
                result, payload = self.run_generator(
                    [patient("u1", "中度患者")], service_start=start, service_end=end,
                )
                self.assertNotEqual(result.returncode, 0)
                self.assertIsNone(payload)
                self.assertIn("服务周期", result.stderr)

    def test_occurrences_respect_patient_windows_across_months(self):
        patients = [
            patient(f"window-{index}", "中度患者", activated=activation)
            for index, activation in enumerate([
                "2026-07-01 00:00:00", "2026-07-31 00:00:00",
                "2026-08-01 12:00:00", "2026-08-15 21:59:58",
            ] * 25)
        ]
        result, payload = self.run_generator(
            patients, service_start="2026-07-31", service_end="2026-08-15",
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        for source, record in zip(patients, payload["records"]):
            occurrence = datetime.fromisoformat(record["occurrenceTime"])
            self.assertGreater(occurrence, datetime.fromisoformat(source["activateTime"]))
            self.assertGreaterEqual(occurrence, datetime(2026, 7, 31))
            self.assertLessEqual(occurrence, datetime(2026, 8, 15, 23, 59, 59))

    def test_one_day_period_includes_last_second(self):
        result, payload = self.run_generator(
            [patient("last-second", "重度患者", activated="2026-08-15 21:59:58")],
            service_start="2026-08-15", service_end="2026-08-15",
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(payload["records"][0]["occurrenceTime"], "2026-08-15 21:59:59")

    def test_rejects_patients_without_available_time_in_service_period(self):
        for activation in ["2026-08-15 23:59:59", "2026-08-16 00:00:00"]:
            with self.subTest(activation=activation):
                result, payload = self.run_generator(
                    [patient("valid-user", "轻度患者", activated="2026-08-01 10:00:00"),
                     patient("blocked-user", "中度患者", activated=activation)],
                    service_start="2026-08-01", service_end="2026-08-15",
                )
                self.assertNotEqual(result.returncode, 0)
                self.assertIsNone(payload)
                for expected in ["blocked-user", "服务周期", "激活时间", "2026-08-15"]:
                    self.assertIn(expected, result.stderr)

    def test_rejects_missing_or_invalid_activation(self):
        for activation in ["", "invalid", "2026-02-30 10:00:00", None]:
            with self.subTest(activation=activation):
                result, payload = self.run_generator([patient("bad-activation", "中度患者", activated=activation)])
                self.assertNotEqual(result.returncode, 0)
                self.assertIsNone(payload)
                self.assertIn("bad-activation", result.stderr)
                self.assertIn("激活时间", result.stderr)

    def test_uses_only_allowed_discovery_methods_and_blank_followup(self):
        result, payload = self.run_generator(
            [patient("u1", "中度患者"), patient("u2", "重度患者")]
        )

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(
            all(
                item["discoveryMethod"] in {"AI用药随访发现", "患者自评反馈"}
                for item in payload["records"]
            )
        )
        self.assertTrue(all(item["followupRecord"] == "" for item in payload["records"]))

    def test_records_include_nonempty_required_structured_fields(self):
        result, payload = self.run_generator([patient("u1", "中度患者")])

        self.assertEqual(result.returncode, 0, result.stderr)
        required = {
            "userid",
            "symptomDescription",
            "severityGrade",
            "treatmentMeasures",
            "treatmentOutcome",
            "remark",
        }
        record = payload["records"][0]
        self.assertTrue(required.issubset(record))
        self.assertTrue(all(record[field] for field in required))

    def test_symptom_description_includes_disease_and_exact_age(self):
        patients = [
            patient("brain-58", "中度患者", disease="脑梗死", age=58),
            patient("heart-76", "重度患者", disease="冠心病心绞痛", age=76),
        ]

        result, payload = self.run_generator(patients)

        self.assertEqual(result.returncode, 0, result.stderr)
        for source, record in zip(patients, payload["records"]):
            self.assertIn(source["disease"], record["symptomDescription"])
            self.assertIn(f"{source['age']}岁", record["symptomDescription"])

    def test_symptom_description_uses_age_band_context(self):
        patients = [
            patient("age-58", "中度患者", age=58),
            patient("age-68", "中度患者", age=68),
            patient("age-78", "中度患者", age=78),
        ]

        result, payload = self.run_generator(patients)

        self.assertEqual(result.returncode, 0, result.stderr)
        descriptions = [item["symptomDescription"] for item in payload["records"]]
        self.assertIn("中年阶段", descriptions[0])
        self.assertIn("老年阶段", descriptions[1])
        self.assertIn("高龄阶段", descriptions[2])

    def test_symptom_descriptions_are_varied_within_same_disease_and_age(self):
        patients = [
            patient(f"same-cohort-{index:02d}", "中度患者", disease="脑梗死", age=66)
            for index in range(24)
        ]

        result, payload = self.run_generator(patients)

        self.assertEqual(result.returncode, 0, result.stderr)
        descriptions = {item["symptomDescription"] for item in payload["records"]}
        self.assertGreaterEqual(len(descriptions), 12)

    def test_content_is_product_and_disease_sensitive_without_claiming_outcome(self):
        result, payload = self.run_generator(
            [
                patient("u1", "中度患者", disease="脑梗死"),
                patient("u2", "中度患者", disease="冠心病心绞痛"),
            ]
        )

        self.assertEqual(result.returncode, 0, result.stderr)
        first, second = payload["records"]
        self.assertNotEqual(first["symptomDescription"], second["symptomDescription"])
        self.assertNotEqual(first["treatmentMeasures"], second["treatmentMeasures"])
        self.assertNotEqual(first["treatmentOutcome"], second["treatmentOutcome"])
        for record in payload["records"]:
            self.assertIn("血栓通胶囊", record["symptomDescription"])
            self.assertIn("血栓通胶囊", record["medicationRelationship"])
            self.assertIn("人工核实", record["medicationRelationship"])
            self.assertNotIn("结构化草案：", record["symptomDescription"])
            self.assertNotIn("人工审核草案：", record["remark"])
            self.assertNotIn("草案", record["symptomDescription"])
            self.assertNotIn("草案", record["remark"])
            self.assertNotRegex(record["treatmentMeasures"], r"停药|减量|加用|继续原方案")
            self.assertRegex(record["treatmentOutcome"], r"待.*核实|未提供")
            self.assertNotRegex(record["treatmentOutcome"], r"痊愈|好转|恢复正常")
            self.assertLessEqual(len(record["treatmentOutcome"]), 160)

        for record in (first, second):
            symptom_tail = record["symptomDescription"].split("反馈可能出现", 1)[1]
            selected_symptoms = "，".join(symptom_tail.split("，")[:2])
            self.assertIn(selected_symptoms, record["treatmentMeasures"])
        self.assertIn("血栓通胶囊", first["treatmentOutcome"])
        self.assertIn("血栓通胶囊", second["treatmentOutcome"])

    def test_notes_include_product_age_sex_disease_allergy_without_draft_prefix(self):
        result, payload = self.run_generator(
            [
                patient(
                    "u1",
                    "重度患者",
                    disease="冠心病心绞痛",
                    gender="女",
                    age=72,
                    allergy="青霉素过敏",
                )
            ]
        )

        self.assertEqual(result.returncode, 0, result.stderr)
        remark = payload["records"][0]["remark"]
        for expected in ["血栓通胶囊", "女", "72岁", "冠心病心绞痛", "青霉素过敏"]:
            self.assertIn(expected, remark)
        self.assertNotIn("草案", remark)

    def test_rejects_input_without_any_target_tag(self):
        result, payload = self.run_generator(
            [patient("u0", "无"), patient("u1", "未知标签")]
        )

        self.assertNotEqual(result.returncode, 0)
        self.assertIsNone(payload)
        self.assertIn("患者标签中没有轻度患者、中度患者或重度患者", result.stderr)


@unittest.skipUnless(NODE and NODE_MODULES, "需要CODEX_NODE和CODEX_NODE_MODULES运行工作簿测试")
class AdverseReactionWorkbookTest(unittest.TestCase):
    def test_builds_and_independently_verifies_workbook(self):
        source_patients = [
            patient("u1", "中度患者", disease="脑梗死"),
            patient("u2", "重度患者", disease="冠心病心绞痛", gender="女", age=72),
        ]
        records = [
            {
                "userid": "u1",
                "disease": "脑梗死",
                "occurrenceTime": "2026-04-11 07:30:00",
                "discoveryMethod": "AI用药随访发现",
                "symptomDescription": "患者在使用血栓通胶囊期间反馈可能出现头晕或乏力，具体情况需人工核实。",
                "severityGrade": "中度",
                "medicationRelationship": "上述表现与血栓通胶囊存在时间关联的可能性，具体因果关系需人工核实。",
                "treatmentMeasures": "建议人工复核症状和当前用药，必要时联系医师，不自行调整用药。",
                "treatmentOutcome": "当前资料未提供处理后转归，需在后续随访中核实并记录。",
                "manualIntervention": "否",
                "followupRecord": "",
                "remark": "涉及血栓通胶囊的相关信息需复核，不构成诊断、处方调整或疗效结论。",
            },
            {
                "userid": "u2",
                "disease": "冠心病心绞痛",
                "occurrenceTime": "2026-04-15 21:59:59",
                "discoveryMethod": "患者自评反馈",
                "symptomDescription": "患者在使用血栓通胶囊期间反馈可能出现明显乏力或胃部不适，具体情况需人工核实。",
                "severityGrade": "重度",
                "medicationRelationship": "上述表现与血栓通胶囊存在时间关联的可能性，具体因果关系需人工核实。",
                "treatmentMeasures": "建议尽快人工干预并复核当前用药，出现紧急情况及时就医，不自行调整用药。",
                "treatmentOutcome": "当前资料未提供处理后转归，需在后续随访中核实并记录。",
                "manualIntervention": "是",
                "followupRecord": "",
                "remark": "涉及血栓通胶囊的相关信息需复核，不构成诊断、处方调整或疗效结论。",
            },
        ]
        payload = {
            "meta": {
                "source": "/tmp/月度患者清单.xlsx",
                "sourceTitle": "测试清单",
                "sourcePatientCount": 2,
                "targetPatientCount": 2,
                "targetTags": ["中度患者", "重度患者"],
                "productName": "血栓通胶囊",
                "servicePeriod": {"start": "2026-04-11", "end": "2026-04-15"},
            },
            "sourcePatients": source_patients,
            "records": records,
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            payload_path = temp / "payload.json"
            workbook_path = temp / "不良反应清单.xlsx"
            preview_dir = temp / "previews"
            report_path = temp / "verification.json"
            payload_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            env = {**os.environ, "CODEX_NODE_MODULES": NODE_MODULES}

            build = subprocess.run(
                [
                    NODE,
                    str(BUILDER),
                    "--payload",
                    str(payload_path),
                    "--template",
                    str(TEMPLATE),
                    "--output",
                    str(workbook_path),
                    "--preview-dir",
                    str(preview_dir),
                ],
                env=env,
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(build.returncode, 0, build.stderr)
            self.assertTrue(workbook_path.is_file())
            rendered = load_workbook(workbook_path, read_only=True, data_only=True)
            sheet = rendered.worksheets[0]
            self.assertEqual(sheet["A1"].value, "不良反应（AE）记录清单")
            self.assertNotIn("草案", sheet["E3"].value)
            self.assertNotIn("草案", sheet["K3"].value)
            headers = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
            self.assertEqual(len(headers), 11)
            self.assertNotIn("发现途径", headers)
            self.assertNotIn("关联随访记录", headers)
            self.assertEqual(sheet["F3"].value, "中度")
            self.assertEqual(sheet["F4"].value, "重度")
            rendered.close()
            self.assertEqual(
                {path.name for path in preview_dir.glob("*.png")},
                {"adverse-first.png", "adverse-middle.png", "adverse-last.png"},
            )

            verify = subprocess.run(
                [
                    NODE,
                    str(VERIFIER),
                    "--payload",
                    str(payload_path),
                    "--workbook",
                    str(workbook_path),
                    "--report",
                    str(report_path),
                ],
                env=env,
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(verify.returncode, 0, verify.stderr)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "passed")
            self.assertEqual(report["rowCount"], 2)
            self.assertEqual(report["distinctUseridCount"], 2)
            self.assertTrue(report["exactUseridOrderMatch"])
            self.assertTrue(report["occurrenceTimesFollowActivation"])
            self.assertTrue(report["occurrenceTimesWithinServicePeriod"])
            self.assertEqual(report["servicePeriod"], payload["meta"]["servicePeriod"])
            self.assertTrue(report["formulaErrors"].endswith("matched 0 entries"))

            invalid_cases = [
                ("before-activation", "2026-04-09 09:00:00", None, "未晚于激活时间"),
                ("equal-activation", "2026-04-10 10:00:00", None, "未晚于激活时间"),
                ("before-period", "2026-04-10 23:59:59", None, "服务周期"),
                ("after-period", "2026-04-16 00:00:00", None, "服务周期"),
                ("before-daily-window", "2026-04-11 07:29:59", None, "每日07:30"),
                ("after-daily-window", "2026-04-11 22:00:00", None, "每日07:30"),
                ("invalid-time", "2026-04-31 09:00:00", None, "日期时间无效"),
                ("missing-period", None, {}, "服务周期"),
                ("invalid-period", None, {"start": "2026-04-11", "end": "2026-04-31"}, "服务周期"),
                ("reversed-period", None, {"start": "2026-04-16", "end": "2026-04-15"}, "服务周期"),
            ]
            for label, occurrence, period, error in invalid_cases:
                with self.subTest(case=label):
                    invalid_payload = copy.deepcopy(payload)
                    if occurrence is not None:
                        invalid_payload["records"][0]["occurrenceTime"] = occurrence
                    if period is not None:
                        invalid_payload["meta"]["servicePeriod"] = period
                    invalid_path = temp / f"{label}.json"
                    invalid_path.write_text(json.dumps(invalid_payload, ensure_ascii=False), encoding="utf-8")
                    rejected_output = temp / f"{label}-build.xlsx"
                    rejected_build = subprocess.run(
                        [NODE, str(BUILDER), "--payload", str(invalid_path),
                         "--template", str(TEMPLATE), "--output", str(rejected_output),
                         "--preview-dir", str(preview_dir)],
                        env=env, capture_output=True, text=True, check=False,
                    )
                    self.assertNotEqual(rejected_build.returncode, 0)
                    self.assertIn(error, rejected_build.stderr)
                    self.assertFalse(rejected_output.exists())

                    tampered_path = temp / f"{label}-tampered.xlsx"
                    tampered = load_workbook(workbook_path)
                    if occurrence is not None:
                        tampered.worksheets[0]["D3"] = occurrence
                    tampered.save(tampered_path)
                    tampered.close()
                    rejected_report = temp / f"{label}-report.json"
                    rejected_verify = subprocess.run(
                        [NODE, str(VERIFIER), "--payload", str(invalid_path),
                         "--workbook", str(tampered_path), "--report", str(rejected_report)],
                        env=env, capture_output=True, text=True, check=False,
                    )
                    self.assertNotEqual(rejected_verify.returncode, 0)
                    self.assertIn(error, rejected_verify.stderr)
                    self.assertFalse(rejected_report.exists())

            tampered = load_workbook(workbook_path)
            tampered.worksheets[0]["D3"] = "2026-04-12 12:00:00"
            tampered_path = temp / "mismatched-time.xlsx"
            tampered.save(tampered_path)
            tampered.close()
            mismatch = subprocess.run(
                [NODE, str(VERIFIER), "--payload", str(payload_path),
                 "--workbook", str(tampered_path), "--report", str(temp / "mismatch-report.json")],
                env=env, capture_output=True, text=True, check=False,
            )
            self.assertNotEqual(mismatch.returncode, 0)
            self.assertIn("发生时间与payload不一致", mismatch.stderr)


if __name__ == "__main__":
    unittest.main()
