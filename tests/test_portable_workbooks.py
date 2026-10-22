"""Public CLI integration without the spreadsheet plugin or artifact-tool."""
import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from PIL import Image

from test_disease_specific_plans import disease_plan, profile
from test_extract_patients import write_workbook


SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / "scripts"))
from extract_patients import EXPECTED_HEADERS

NODE = os.environ.get("NODE_BINARY") or shutil.which("node")


@unittest.skipUnless(NODE, "需要Node.js")
class PortableWorkbookTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.env = {**os.environ, "PYTHON_BINARY": sys.executable, "AUTO_MEDICATION_SEARCH": "0",
                    "CODEX_NODE_MODULES": str(self.root / "missing-node-modules")}
        self.env.pop("WORKBOOK_BACKEND", None)

    def run_script(self, name, *args, success=True):
        binary = NODE if name.endswith(".mjs") else sys.executable
        result = subprocess.run(
            [binary, str(SKILL / "scripts" / name), *map(str, args)],
            env=self.env, capture_output=True, text=True, encoding="utf-8",
        )
        if success:
            self.assertEqual(result.returncode, 0, result.stderr[-5000:])
        else:
            self.assertNotEqual(result.returncode, 0)
        return result

    def write_json(self, name, data):
        target = self.root / name
        target.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return target

    def monthly_payload(self):
        headers = EXPECTED_HEADERS[:10] + ["关联厂家用药"] + EXPECTED_HEADERS[10:]
        rows = [[index, uid, "=1+1" if index == 1 else "测试患者", "2026-10-08 15:16:07",
                 "女", 42, "", "", "测试疾病", "无", "  测试产品  ",
                 0, 0, 0, 0, "", "否", "", "轻度患者"]
                for index, uid in enumerate(["00009", "u2", "u1"], start=1)]
        source = self.root / "月度患者清单19列.xlsx"
        # Input names that resemble formulas are explicitly stored as text.
        write_workbook(source, headers, rows)
        book = openpyxl.load_workbook(source)
        book.active["C3"].data_type = "s"
        book.save(source)
        book.close()
        before = source.read_bytes()
        patients = self.root / "patients.json"
        self.run_script("extract_patients.py", "--source", source, "--output", patients)
        self.assertEqual(source.read_bytes(), before)
        spec = profile([disease_plan("测试方案", ["测试疾病"], [])])
        product = self.write_json("profile.json", spec)
        payload = self.root / "payload.json"
        self.run_script("generate_payload.py", "--patients", patients, "--profile", product,
                        "--output", payload)
        return payload

    def check_previews(self, directory, prefixes):
        for prefix in prefixes:
            for segment in ["first", "middle", "last"]:
                target = directory / f"{prefix}-{segment}.png"
                self.assertTrue(target.is_file(), str(target))
                with Image.open(target) as image:
                    self.assertGreater(image.width, 100)
                    self.assertGreater(image.height, 50)
                    self.assertIsNotNone(image.convert("RGB").getextrema())

    def test_monthly_cli_without_plugin_preserves_data_and_rejects_saved_tampering(self):
        payload_path = self.monthly_payload()
        payload = json.loads(payload_path.read_text())
        self.run_script("build_workbooks.mjs", "--payload", payload_path,
                        "--reminder-template", SKILL / "assets/medication-reminder-template.xlsx",
                        "--medication-template", SKILL / "assets/medication-list-template.xlsx",
                        "--output-dir", self.root)
        reminder = self.root / "用药提醒_测试产品.xlsx"
        medication = self.root / "用药方案_测试产品.xlsx"
        report = self.root / "verification.json"
        args = ["--payload", payload_path, "--reminder", reminder,
                "--medication", medication, "--report", report]
        self.run_script("verify_workbooks.mjs", *args)
        self.assertEqual(json.loads(report.read_text())["status"], "passed")
        self.check_previews(self.root, ["reminder", "medication"])
        book = openpyxl.load_workbook(reminder)
        sheet = book.active
        self.assertEqual((sheet.max_row, sheet.max_column), (5, 18))
        self.assertEqual(sheet["M3"].value, "  测试产品  ")
        self.assertEqual(sheet["C3"].value, "=1+1")
        self.assertEqual(sheet["C3"].data_type, "s")
        self.assertEqual([sheet.cell(r, 2).value for r in range(3, 6)], ["00009", "u2", "u1"])
        self.assertEqual(sheet.freeze_panes, "A3")
        self.assertIn("A1:R1", {str(value) for value in sheet.merged_cells.ranges})
        self.assertEqual(next(iter(sheet.tables.values())).ref, "A2:R5")
        book.close()
        for address, value in [("I3", "错误剂量"), ("K3", "08:00（餐后）"),
                               ("M3", "缺少四项"), ("G3", "其他药"), ("I3", "=1/0")]:
            with self.subTest(address=address, value=value):
                saved = openpyxl.load_workbook(medication)
                saved.active[address] = value
                tampered = self.root / "tampered.xlsx"
                saved.save(tampered)
                saved.close()
                bad_report = self.root / "bad-report.json"
                self.run_script("verify_workbooks.mjs", "--payload", payload_path,
                                "--reminder", reminder, "--medication", tampered,
                                "--report", bad_report, success=False)
                self.assertFalse(bad_report.exists())
        invalid = copy.deepcopy(payload)
        invalid["medicationItems"][0]["precautions"] = "缺少四项"
        invalid_path = self.write_json("invalid.json", invalid)
        rejected = self.root / "rejected"
        self.run_script("build_workbooks.mjs", "--payload", invalid_path,
                        "--reminder-template", SKILL / "assets/medication-reminder-template.xlsx",
                        "--medication-template", SKILL / "assets/medication-list-template.xlsx",
                        "--output-dir", rejected, success=False)
        self.assertFalse(list(rejected.glob("*.xlsx")))
        if os.environ.get("WORKBOOK_TEST_ARTIFACTS"):
            destination = Path(os.environ["WORKBOOK_TEST_ARTIFACTS"]) / "portable"
            destination.mkdir(parents=True, exist_ok=True)
            for target in [reminder, medication, *self.root.glob("*.png")]:
                shutil.copy2(target, destination / target.name)

    def test_adverse_cli_without_plugin_keeps_occurrence_and_wording_validation(self):
        self.monthly_payload()
        payload = self.root / "adverse.json"
        self.run_script("generate_adverse_reactions.py", "--patients", self.root / "patients.json",
                        "--product", "测试产品", "--service-start", "2026-10-01",
                        "--service-end", "2026-10-31", "--output", payload)
        output = self.root / "不良反应.xlsx"
        preview = self.root / "previews"
        self.run_script("build_adverse_reaction_workbook.mjs", "--payload", payload,
                        "--template", SKILL / "assets/adverse-reaction-template.xlsx",
                        "--output", output, "--preview-dir", preview)
        report = self.root / "adverse-verification.json"
        self.run_script("verify_adverse_reaction_workbook.mjs", "--payload", payload,
                        "--workbook", output, "--report", report)
        self.assertTrue(json.loads(report.read_text())["occurrenceTimesMatchActivationPeriodRule"])
        self.check_previews(preview, ["adverse"])
        if os.environ.get("WORKBOOK_TEST_ARTIFACTS"):
            destination = Path(os.environ["WORKBOOK_TEST_ARTIFACTS"]) / "portable-adverse"
            destination.mkdir(parents=True, exist_ok=True)
            for target in [output, *preview.glob("*.png")]:
                shutil.copy2(target, destination / target.name)
        book = openpyxl.load_workbook(output)
        book.active["D3"] = "2026-10-08 09:00:00"
        book.save(output)
        book.close()
        self.run_script("verify_adverse_reaction_workbook.mjs", "--payload", payload,
                        "--workbook", output, "--report", self.root / "rejected.json",
                        success=False)
        self.assertFalse((self.root / "rejected.json").exists())

    def test_environment_check_reports_backend_and_actionable_missing_python(self):
        result = self.run_script("check_workbook_environment.mjs")
        report = json.loads(result.stdout)
        self.assertEqual(report["backend"], "portable")
        self.assertTrue(report["ready"])
        self.env["PYTHON_BINARY"] = str(self.root / "不存在的Python.exe")
        result = self.run_script("check_workbook_environment.mjs", success=False)
        self.assertIn("PYTHON_BINARY", result.stderr)

    def test_portable_monthly_handles_3255_patients_without_losing_order(self):
        payload_path = self.monthly_payload()
        data = json.loads(payload_path.read_text())
        original = copy.deepcopy(data)
        data["patients"], data["records"], data["medicationItems"] = [], [], []
        per_user_meta = [key for key, value in data["meta"].items()
                         if key.endswith("ByUserid") and isinstance(value, dict)]
        for key in per_user_meta:
            data["meta"][key] = {}
        for index in range(3255):
            uid = f"{index:032d}"
            for collection in ["patients", "records", "medicationItems"]:
                entry = copy.deepcopy(original[collection][0])
                entry["userid"] = uid
                data[collection].append(entry)
            for key in per_user_meta:
                values = original["meta"][key]
                data["meta"][key][uid] = copy.deepcopy(values[original["patients"][0]["userid"]])
        data["meta"]["patientCount"] = 3255
        data["meta"]["minimumUniqueMedicationPlanCount"] = 33
        data["meta"]["uniqueMedicationPlanTargetMet"] = False
        data["meta"]["uniqueMedicationPlanShortfall"] = 32
        payload_path = self.write_json("large.json", data)
        self.run_script("build_workbooks.mjs", "--payload", payload_path,
                        "--reminder-template", SKILL / "assets/medication-reminder-template.xlsx",
                        "--medication-template", SKILL / "assets/medication-list-template.xlsx",
                        "--output-dir", self.root)
        report = self.root / "large-verification.json"
        self.run_script("verify_workbooks.mjs", "--payload", payload_path,
                        "--reminder", self.root / "用药提醒_测试产品.xlsx",
                        "--medication", self.root / "用药方案_测试产品.xlsx", "--report", report)
        self.assertEqual(json.loads(report.read_text())["patientCount"], 3255)
