import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from docx import Document
from test_extract_insight_sources import write_book
from extract_insight_sources import build_insight
import generate_insight_charts as chart_module
from generate_insight_charts import generate_charts
from build_insight_report import build_report
from validate_insight_report import validate_report


class MonthlyInsightTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        identity = ['患者唯一标识', '性别', '年龄', '所属地区', '疾病']
        write_book(self.root/'patients.xlsx', identity + ['AI用药提醒次数'],
                   [['001', '女', 60, '江苏', '疾病甲', 12], ['002', '男', 50, '江苏', '疾病乙', 0]],
                   '2026年04月-月度患者服务清单')
        write_book(self.root/'followups.xlsx', identity + ['随访时间'] + [f'{i}、真实问题{i}' for i in range(1, 11)],
                   [['001','女',60,'江苏','疾病甲','2026-04-10']+['A、选项']*10,
                    ['outsider','男',30,'江苏','疾病乙','2026-04-10']+['E、选项']*10])
        write_book(self.root/'symptoms.xlsx', identity + ['自评时间'] + [f'{i}、真实症状{i}' for i in range(1,7)],
                   [['001','女',60,'江苏','疾病甲','2026-04-10']+['B、轻微']*6,
                    ['002','男',50,'江苏','疾病乙','2026-05-01']+['E、严重']*6])
        write_book(self.root/'reminders.xlsx', ['患者唯一标识','性别','年龄','疾病','联合用药','用药方案确认时间','用药方案','用药周期'],
                   [['001','女',60,'疾病甲','产品甲、药品乙','2026-04-09','产品甲1粒、药品乙10mg','21天'],
                    ['002','男',50,'疾病乙','产品甲','2026-05-01','产品甲2粒','30天']])
        write_book(self.root/'ae.xlsx', ['患者ID','不良反应发生时间','不良反应严重程度分级','疾病'],
                   [['001','2026-04-11','重度','疾病甲']])
        self.paths = [str(self.root/name) for name in ['patients.xlsx','followups.xlsx','symptoms.xlsx','reminders.xlsx','ae.xlsx']]

    def insight(self, paths=None, start='2026-04-01', end='2026-04-30'):
        return build_insight(paths or self.paths, '产品甲', start, end)

    def test_five_sources_have_distinct_roles_and_source_backed_metrics(self):
        report = self.insight(); m=report['metrics']
        self.assertEqual(report['metadata']['inputMode'], 'monthly')
        self.assertEqual(m['serviceExecution']['medicationReminders'],12)
        self.assertEqual(m['serviceExecution']['followupRecords'],1)
        self.assertEqual(m['reminderPlanCoverage']['numerator'],1)
        self.assertEqual(m['symptoms']['validAssessments'],1)
        self.assertEqual(m['medications']['drugDistribution'], [{'label':'产品甲','count':1},{'label':'药品乙','count':1}])
        self.assertIsNone(m['healthPlanCoverage'])
        self.assertIsNone(m['serviceExecution']['estimatedResponseRate'])
        self.assertIsNone(m['serviceExecution']['temperatureMonitoring'])
        self.assertEqual(m['medications']['productSpecificationDistribution'],[])
        self.assertEqual(report['sourceDiagnostics']['roles']['followups']['unmatchedRows'],1)

    def test_mismatched_period_does_not_reuse_monthly_summary(self):
        report=self.insight(start='2026-09-01',end='2026-09-30')
        self.assertIsNone(report['metrics']['serviceExecution']['medicationReminders'])
        self.assertEqual(report['metrics']['reminderPlanCoverage']['numerator'],0)
        self.assertEqual(report['metrics']['followupCoverage']['numerator'],0)
        self.assertFalse(report['metrics']['monthlySummary']['matchesPeriod'])

    def test_missing_and_duplicate_required_roles_stop(self):
        with self.assertRaisesRegex(ValueError,'缺少工作簿角色'):
            self.insight(paths=self.paths[1:])
        duplicate=self.root/'duplicate.xlsx';duplicate.write_bytes((self.root/'reminders.xlsx').read_bytes())
        with self.assertRaisesRegex(ValueError,'角色重复'):
            self.insight(paths=self.paths[:-1]+[str(duplicate)])

    def test_partial_month_and_missing_counts_remain_unavailable(self):
        report = self.insight(start='2026-04-05', end='2026-04-30')
        self.assertIsNone(report['metrics']['serviceExecution']['medicationReminders'])
        path = self.root/'patients.xlsx'
        book = load_workbook(path); book.active.cell(3, 6).value = None; book.save(path); book.close()
        report = self.insight()
        self.assertTrue(report['metrics']['monthlySummary']['matchesPeriod'])
        self.assertIsNone(report['metrics']['serviceExecution']['medicationReminders'])

    def test_unknown_age_and_annotated_severity(self):
        path = self.root/'patients.xlsx'
        book = load_workbook(path); book.active.cell(3, 3).value = None; book.save(path); book.close()
        path = self.root/'ae.xlsx'
        book = load_workbook(path); book.active.cell(2, 3).value = '重度（3级）'; book.save(path); book.close()
        metrics = self.insight()['metrics']
        unknown = next(x for x in metrics['moduleCoverageByAge'] if x['label']=='无记录')
        self.assertEqual(unknown['patientCount'], 1)
        self.assertEqual(unknown['modules']['智能随访']['numerator'], 1)
        self.assertIn({'label':'无记录','count':1}, metrics['ageDistribution'])
        self.assertEqual(next(x['count'] for x in metrics['riskDistribution'] if x['label']=='高风险'), 1)

    def test_unanswered_question_is_not_plotted_as_zero_percent(self):
        path = self.root/'followups.xlsx'
        book = load_workbook(path); book.active.cell(2, 12).value = None; book.save(path); book.close()
        report = self.insight()
        self.assertIsNone(report['metrics']['followups']['questions'][5]['positiveRate']['value'])
        with patch.object(chart_module, '_bar_chart', wraps=chart_module._bar_chart) as plot:
            generate_charts(report, self.root/'missing-answer-charts')
        call = next(call for call in plot.call_args_list if call.args[0].name == '08-followup-positive-rate.png')
        self.assertNotIn('Q6', call.args[1])
        self.assertIn('Q5', call.args[1])

    def test_full_pipeline_and_empty_period_omit_unavailable_statistics(self):
        for paths,start,end in [(self.paths,'2026-04-01','2026-04-30'),
                                (self.paths[:-1],'2026-04-01','2026-04-30'),
                                (self.paths,'2026-09-01','2026-09-30')]:
            with self.subTest(start=start,files=len(paths)):
                insight=self.insight(paths,start,end)
                payload=self.root/'insight.json';payload.write_text(json.dumps(insight,ensure_ascii=False))
                manifest=generate_charts(insight,self.root/'charts')
                output=self.root/'report.docx';build_report(payload,manifest,None,output)
                self.assertEqual(validate_report(output,payload),[])
                doc=Document(output);text='\n'.join(p.text for p in doc.paragraphs)
                cells='\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
                self.assertNotIn('健康管理方案覆盖',cells)
                self.assertNotIn('估算响应率',cells)
                self.assertNotIn('产品规格',cells)
                if start=='2026-04-01': self.assertIn('真实问题6',text)
                if len(paths)==4:
                    self.assertNotIn('患者发生率',cells)
                    self.assertIsNone(insight['metrics']['adverseEventRate'])
                doc.add_paragraph('健康管理方案覆盖率100.0%，估算响应率100.0%。')
                doc.save(output)
                self.assertTrue(validate_report(output,payload))
