import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from validate_insight_report import validate_report


class ValidateInsightReportTest(unittest.TestCase):
    def test_requires_cover_toc_page_number_and_forbids_chart_explanation(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            docx_path = root / "report.docx"
            insight_path = root / "insight.json"
            document = Document()
            document.add_paragraph("患者洞察报告")
            document.add_paragraph("图表说明：不应出现在最终报告")
            document.save(docx_path)
            insight_path.write_text(json.dumps({
                "metadata": {
                    "product": "产品",
                    "period": {"start": "2026-07-01", "end": "2026-07-31"},
                    "patientCount": 1,
                },
                "metrics": {"serviceExecution": {"totalPrompts": 1}},
            }, ensure_ascii=False), encoding="utf-8")

            errors = validate_report(docx_path, insight_path)
            self.assertIn("缺少患者洞察报告封面或封面表格", errors)
            self.assertIn("缺少动态目录字段", errors)
            self.assertIn("正文缺少PAGE页码字段", errors)
            self.assertIn("正文页码未从1开始", errors)
            self.assertIn("正文不得包含图表说明段", errors)


if __name__ == "__main__":
    unittest.main()
