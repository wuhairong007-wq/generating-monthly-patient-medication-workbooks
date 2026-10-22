"""Static workbook I/O and previews; clinical and row validation remains in JS."""
import json
import math
import os
import sys
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from PIL import Image, ImageDraw, ImageFont
except ImportError as exc:
    requirements = Path(__file__).resolve().parents[1] / "requirements-workbooks.txt"
    raise SystemExit(
        f"缺少表格后端依赖：{exc}。请在当前Python环境执行：\n"
        f'"{sys.executable}" -m pip install -r "{requirements}"'
    ) from exc


def read_workbook(target):
    book = openpyxl.load_workbook(target, data_only=False)
    sheets = []
    for sheet in book:
        values = [[value.isoformat(sep=" ") if isinstance(value, datetime)
                   else value.isoformat() if isinstance(value, date) else value
                   for value in row] for row in sheet.values]
        while values and all(value is None for value in values[-1]):
            values.pop()
        used_columns = max((index + 1 for row in values for index, value in enumerate(row)
                            if value is not None), default=0)
        values = [row[:used_columns] for row in values]
        sheets.append({
            "name": sheet.title, "values": values,
            "tables": [{"name": table.name, "ref": table.ref} for table in sheet.tables.values()],
            "formulas": [cell.coordinate for row in sheet for cell in row if cell.data_type == "f"],
        })
    book.close()
    return {"sheets": sheets}


def preview_font():
    windows = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    candidates = [
        os.environ.get("WORKBOOK_PREVIEW_FONT"),
        str(windows / "msyh.ttc"), str(windows / "simhei.ttf"),
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return ImageFont.truetype(candidate, 15)
    raise ValueError("缺少中文预览字体；安装中文字体或设置WORKBOOK_PREVIEW_FONT为字体文件路径。")


def text_lines(font):
    @lru_cache(maxsize=20000)
    def wrap(text, width):
        lines = []
        for paragraph in text.split("\n"):
            line = ""
            for char in paragraph:
                if line and font.getlength(line + char) > width:
                    lines.append(line)
                    line = ""
                line += char
            lines.append(line)
        return lines or [""]
    return wrap


def render_preview(sheet, rows, widths, font, wrap, target):
    heights = [math.ceil((sheet.row_dimensions[row].height or 30) * 4 / 3) for row in rows]
    image = Image.new("RGB", (sum(widths) + 1, sum(heights) + 1), "white")
    draw = ImageDraw.Draw(image)
    y = 0
    for row, height in zip(rows, heights):
        if row == 1:
            draw.rectangle((0, y, sum(widths), y + height), fill="#FFFFFF", outline="#D9D9D9")
            draw.text((8, y + 6), str(sheet["A1"].value or ""), fill="#156082", font=font)
        else:
            x = 0
            for column, width in enumerate(widths, start=1):
                fill = "#156082" if row == 2 else "#E9F3F7" if row % 2 else "#FFFFFF"
                draw.rectangle((x, y, x + width, y + height), fill=fill, outline="#D9D9D9")
                cell = sheet.cell(row, column)
                text = str(cell.value if cell.value is not None else "")
                lines = wrap(text, width - 12)
                for line_index, line in enumerate(lines):
                    text_x = x + (width - font.getlength(line)) / 2 if cell.alignment.horizontal == "center" else x + 6
                    draw.text((text_x, y + 5 + line_index * 21), line,
                              fill="#FFFFFF" if row == 2 else "#202020", font=font)
                x += width
        y += height
    target.parent.mkdir(parents=True, exist_ok=True)
    image.save(target)


def build_workbook(spec):
    template = Path(spec["template"]).resolve()
    output = Path(spec["output"]).resolve()
    if template == output:
        raise ValueError("输出不能覆盖模板")
    book = openpyxl.load_workbook(template)
    sheet = book.worksheets[0]
    headers = spec["headers"]
    actual_headers = [sheet.cell(2, column).value for column in range(1, sheet.max_column + 1)]
    while actual_headers and actual_headers[-1] is None:
        actual_headers.pop()
    if actual_headers != headers:
        raise ValueError("模板表头不匹配")
    font = preview_font()
    wrap = text_lines(font)
    column_widths = list(spec["widths"])
    for column in spec.get("textColumns", []):
        maximum = max((font.getlength(str(row[column - 1])) for row in spec["rows"]), default=0)
        column_widths[column - 1] = max(column_widths[column - 1], math.ceil((maximum + 12 - 5) / 7))
    widths = [int(width * 7 + 5) for width in column_widths]
    # This fallback handles the skill's static lists, never pretends to calculate formulas.
    for other in book:
        if other != sheet and any(cell.value is not None for row in other for cell in row):
            raise ValueError("备用模板含额外数据工作表，需先核对模板")
    for name in list(sheet.tables):
        del sheet.tables[name]
    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))
    if sheet.max_row >= 3:
        sheet.delete_rows(3, sheet.max_row - 2)
    if sheet.max_column > len(headers):
        sheet.delete_cols(len(headers) + 1, sheet.max_column - len(headers))
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet["A1"] = spec["title"]
    sheet["A1"].data_type = "s"
    sheet["A1"].font = Font(name=font.getname()[0], size=16, bold=True, color="156082")
    sheet.row_dimensions[1].height = 30
    edge = Side(style="thin", color="D9D9D9")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        cell.font = Font(name=font.getname()[0], size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="156082")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        sheet.column_dimensions[get_column_letter(column)].width = column_widths[column - 1]
    sheet.row_dimensions[2].height = max(
        32, max(len(wrap(header, widths[i] - 12)) for i, header in enumerate(headers)) * 16 + 8
    )
    for row_index, values in enumerate(spec["rows"], start=3):
        required_height = spec.get("rowHeight", 72)
        if len(values) != len(headers):
            raise ValueError("写入行宽与表头不一致")
        for column, value in enumerate(values, start=1):
            if isinstance(value, str) and len(value) > 32767:
                raise ValueError(f"第{row_index}行第{column}列超出Excel单元格长度，不能截断")
            cell = sheet.cell(row_index, column, value)
            if isinstance(value, str):
                cell.data_type = "s"
            cell.font = Font(name=font.getname()[0], size=11)
            cell.fill = PatternFill("solid", fgColor="E9F3F7" if row_index % 2 else "FFFFFF")
            cell.alignment = Alignment(
                horizontal="center" if column in spec.get("centerColumns", []) else "left",
                vertical="top", wrap_text=True,
            )
            cell.border = border
            if column == 2 or column in spec.get("textColumns", []):
                cell.number_format = "@"
            line_count = len(wrap(str(value if value is not None else ""), widths[column - 1] - 12))
            required_height = max(required_height, math.ceil((line_count * 21 + 12) * 3 / 4))
        if required_height > 409:
            raise ValueError(f"第{row_index}行文字超出Excel最大行高，需调整列宽后重试，不能截断")
        sheet.row_dimensions[row_index].height = required_height
    last_row = len(spec["rows"]) + 2
    table = Table(displayName=spec["tableName"], ref=f"A2:{get_column_letter(len(headers))}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    book.close()
    # Render saved values and dimensions so previews reflect the delivered file.
    saved = openpyxl.load_workbook(output, data_only=False)
    sheet = saved.worksheets[0]
    middle = max(3, (3 + last_row) // 2)
    segments = {
        "first": range(1, min(last_row, 8) + 1),
        "middle": range(max(3, middle - 3), min(last_row, middle + 3) + 1),
        "last": range(max(3, last_row - 6), last_row + 1),
    }
    for name, rows in segments.items():
        render_preview(sheet, rows, widths, font, wrap,
                       Path(spec["previewDir"]) / f'{spec["prefix"]}-{name}.png')
    saved.close()
    return {"status": "passed", "backend": "portable", "output": str(output),
            "rowCount": len(spec["rows"]), "previewsRendered": True, "previewRenderer": "Pillow"}


def main():
    request = json.load(sys.stdin)
    if sys.argv[1] == "check":
        if sys.version_info < (3, 10):
            raise ValueError("需要Python 3.10或更高版本")
        font = preview_font() if request.get("preview", True) else None
        result = {"ready": True, "python": sys.version.split()[0],
                  "openpyxl": openpyxl.__version__, "font": font.getname()[0] if font else None}
    elif sys.argv[1] == "read":
        result = read_workbook(request["path"])
    elif sys.argv[1] == "build":
        result = build_workbook(request)
    else:
        raise ValueError("不支持的表格操作")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
