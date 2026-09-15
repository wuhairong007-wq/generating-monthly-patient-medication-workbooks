#!/usr/bin/env python3
import argparse
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROLE_RULES = {
    "patients": {"userid", "性别", "年龄", "地区", "疾病"},
    "healthPlans": {"userid", "AI健康管理师介绍", "AI健康管理方案"},
    "tracking": {"患者ID", "体温监测次数", "用药提醒次数", "患者响应率"},
    "followups": {"患者ID", "随访时间"},
    "symptomAssessments": {"患者ID", "自评时间"},
    "medications": {"userid", "用药方案确认时间", "药品名称", "规格"},
    "adverseEvents": {"患者ID", "不良反应发生时间", "不良反应严重程度分级"},
}

REQUIRED_ROLES = set(ROLE_RULES) - {"adverseEvents"}
ROLE_RULES["medicationReminders"] = {"userid", "联合用药", "用药方案确认时间", "用药方案", "用药周期"}


DATE_FIELDS = {
    "followups": "随访时间",
    "symptomAssessments": "自评时间",
    "medications": "用药方案确认时间",
    "medicationReminders": "用药方案确认时间",
    "adverseEvents": "不良反应发生时间",
}


def clean(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value)
    match = re.search(r"(\d{4})[-/]?(\d{2})[-/]?(\d{2})", text)
    if not match:
        return None
    return date(*map(int, match.groups()))


def ratio(numerator, denominator):
    value = numerator / denominator if denominator else None
    return {
        "numerator": numerator,
        "denominator": denominator,
        "value": value,
        "display": f"{value * 100:.1f}%" if value is not None else "无记录",
    }


def distribution(values):
    counts = Counter(clean(value) or "无记录" for value in values)
    return [{"label": label, "count": count} for label, count in counts.most_common()]


def role_candidates(headers):
    # Service workbooks repeat patient demographics; dated/service fields identify their role.
    comparable = set(headers)
    if "userid" in comparable:
        comparable.add("患者ID")
    candidates = [role for role, required in ROLE_RULES.items() if required.issubset(comparable)]
    if len(candidates) > 1 and "patients" in candidates:
        candidates.remove("patients")
    return candidates


def find_header(rows):
    for index, row in enumerate(rows[:8]):
        aliases = {"患者唯一标识": "userid", "所属地区": "地区"}
        headers = [aliases.get(clean(value), clean(value)) for value in row]
        header_set = set(headers)
        candidates = role_candidates(header_set)
        if candidates:
            return index, headers, candidates
    raise ValueError("前8行未识别到受支持的表头")


def detect_role(headers):
    header_set = set(headers)
    candidates = role_candidates(header_set)
    if len(candidates) != 1:
        raise ValueError(f"工作簿角色识别不唯一：{candidates or '无匹配'}")
    return candidates[0]


def read_source(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        raw_rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    header_index, headers, _ = find_header(raw_rows)
    role = detect_role(headers)
    records = []
    for excel_row, values in enumerate(raw_rows[header_index + 1 :], start=header_index + 2):
        if not any(value not in (None, "") for value in values):
            continue
        record = {header: clean(values[index]) if index < len(values) else "" for index, header in enumerate(headers) if header}
        record["_sourceRow"] = excel_row
        if role == "patients":
            record["_sourceTitle"] = " ".join(clean(cell) for row in raw_rows[:header_index] for cell in row if cell)
        records.append(record)
    return role, records, header_index + 1


def normalize_patient_id(record):
    return clean(record.get("userid") or record.get("患者ID"))


def filter_period(role, records, start, end):
    field = DATE_FIELDS.get(role)
    if not field:
        return records
    included = []
    for record in records:
        parsed = parse_date(record.get(field))
        if parsed is None:
            raise ValueError(f"{role} 第{record['_sourceRow']}行的{field}无法解析")
        if start <= parsed <= end:
            included.append(record)
    return included


def parse_number(value):
    match = re.search(r"-?\d+(?:\.\d+)?", clean(value))
    return float(match.group()) if match else 0.0


def symptom_score(value):
    text = clean(value)
    match = re.search(r"([1-5])\s*分", text)
    if match:
        return int(match.group(1))
    letter = text[:1].upper()
    return {"A": 1, "B": 2, "C": 3, "D": 4, "E": 5}.get(letter)


def age_bucket(value):
    age = int(parse_number(value))
    if age <= 0:
        return "无记录"
    if age <= 30:
        return "30岁及以下"
    if age <= 40:
        return "31–40岁"
    if age <= 50:
        return "41–50岁"
    if age <= 60:
        return "51–60岁"
    return "61岁及以上"


def question_fields(records, count):
    if not records:
        return []
    headers = [header for header in records[0] if re.match(r"^\d+[、.]", header)]
    return sorted(headers, key=lambda header: int(re.match(r"^(\d+)", header).group(1)))[:count]


def build_metrics(data, product):
    patients = data["patients"]
    patient_ids = {normalize_patient_id(record) for record in patients if normalize_patient_id(record)}
    patient_count = len(patient_ids)
    health_plan_ids = {normalize_patient_id(record) for record in data["healthPlans"] if normalize_patient_id(record) in patient_ids}
    followup_ids = {normalize_patient_id(record) for record in data["followups"] if normalize_patient_id(record) in patient_ids}
    symptom_ids = {normalize_patient_id(record) for record in data["symptomAssessments"] if normalize_patient_id(record) in patient_ids}

    tracking = [record for record in data["tracking"] if normalize_patient_id(record) in patient_ids]
    reminder_total = int(sum(parse_number(record.get("用药提醒次数")) for record in tracking))
    temperature_total = int(sum(parse_number(record.get("体温监测次数")) for record in tracking))
    vital_total = int(sum(parse_number(record.get("血压、心率监测次数")) for record in tracking))
    weighted_response_numerator = sum(
        (parse_number(record.get("用药提醒次数")) + parse_number(record.get("体温监测次数")) + parse_number(record.get("血压、心率监测次数")))
        * parse_number(record.get("患者响应率"))
        / 100
        for record in tracking
    )
    total_prompts = reminder_total + temperature_total + vital_total

    followup_questions = question_fields(data["followups"], 10)
    followup_dimensions = []
    for field in followup_questions:
        values = [record.get(field) for record in data["followups"] if record.get(field)]
        positive = sum(clean(value)[:1].upper() in {"A", "B"} for value in values)
        followup_dimensions.append({"question": field, "distribution": distribution(values), "positiveRate": ratio(positive, len(values))})

    symptom_questions = question_fields(data["symptomAssessments"], 6)
    dimension_scores = [[] for _ in symptom_questions]
    total_scores = []
    for record in data["symptomAssessments"]:
        scores = [symptom_score(record.get(field)) for field in symptom_questions]
        if len(scores) == 6 and all(score is not None for score in scores):
            total_scores.append(sum(scores))
            for index, score in enumerate(scores):
                dimension_scores[index].append(score)
    dimension_means = [round(sum(scores) / len(scores), 2) if scores else None for scores in dimension_scores]
    score_bands = Counter()
    for score in total_scores:
        if score <= 10:
            score_bands["≤10分"] += 1
        elif score <= 15:
            score_bands["11–15分"] += 1
        elif score <= 20:
            score_bands["16–20分"] += 1
        elif score <= 25:
            score_bands["21–25分"] += 1
        else:
            score_bands[">25分"] += 1

    medications = [record for record in data["medications"] if normalize_patient_id(record) in patient_ids]
    medication_patient_counts = Counter(normalize_patient_id(record) for record in medications)
    product_records = [record for record in medications if clean(record.get("药品名称")) == product]

    adverse_events_provided = "adverseEvents" in data
    adverse_events = [record for record in data.get("adverseEvents", []) if normalize_patient_id(record) in patient_ids]
    adverse_ids = {normalize_patient_id(record) for record in adverse_events}
    severe_ids = {
        normalize_patient_id(record)
        for record in adverse_events
        if re.match(r"^(?:高度|重度|严重|危及生命)(?:$|[\s（(])", clean(record.get("不良反应严重程度分级")))
        or clean(record.get("是否触发人工干预")) == "是"
    }
    high_risk = len(severe_ids)
    medium_risk = len(adverse_ids - severe_ids)
    low_risk = max(patient_count - high_risk - medium_risk, 0)

    sex_counts = Counter(clean(record.get("性别")) or "无记录" for record in patients)
    age_counts = Counter(age_bucket(record.get("年龄")) for record in patients if parse_number(record.get("年龄")) > 0)
    region_counts = Counter(clean(record.get("地区")) or "无记录" for record in patients)
    disease_counts = Counter(clean(record.get("疾病")) or "无记录" for record in patients)

    sex_age_counts = Counter(
        (clean(record.get("性别")) or "无记录", age_bucket(record.get("年龄")))
        for record in patients
        if parse_number(record.get("年龄")) > 0
    )
    sex_age_distribution = [
        {"sex": sex, "age": age, "count": count}
        for (sex, age), count in sorted(sex_age_counts.items(), key=lambda item: (-item[1], item[0]))
    ]
    allergy_counts = Counter(clean(record.get("既往过敏史")) or "无记录" for record in patients)
    tracking_ids = {normalize_patient_id(record) for record in tracking if normalize_patient_id(record) in patient_ids}
    module_sets = {
        "健康管理方案": health_plan_ids,
        "用药提醒": tracking_ids,
        "智能随访": followup_ids,
        "症状自评": symptom_ids,
    }

    def grouped_coverage(field):
        groups = defaultdict(list)
        for record in patients:
            patient_id = normalize_patient_id(record)
            if patient_id:
                key = clean(record.get(field)) if field != "年龄" else age_bucket(record.get("年龄"))
                groups[key or "无记录"].append(patient_id)
        result = []
        for key, ids in sorted(groups.items(), key=lambda item: (-len(item[1]), item[0])):
            result.append({
                "label": key,
                "patientCount": len(ids),
                "modules": {name: ratio(len(set(ids) & covered), len(ids)) for name, covered in module_sets.items()},
            })
        return result

    combination_modes = Counter()
    for patient_id in patient_ids:
        drugs = sorted({clean(record.get("药品名称")) for record in medications if normalize_patient_id(record) == patient_id and clean(record.get("药品名称"))})
        if drugs:
            combination_modes[" + ".join(drugs)] += 1

    disease_score_accumulator = defaultdict(lambda: {"count": 0, "sums": [0] * len(symptom_questions)})
    for record in data["symptomAssessments"]:
        scores = [symptom_score(record.get(field)) for field in symptom_questions]
        disease = clean(record.get("疾病")) or "无记录"
        if len(scores) == 6 and all(score is not None for score in scores):
            entry = disease_score_accumulator[disease]
            entry["count"] += 1
            entry["sums"] = [a + b for a, b in zip(entry["sums"], scores)]
    disease_score_distribution = []
    for disease, entry in sorted(disease_score_accumulator.items(), key=lambda item: (-item[1]["count"], item[0])):
        disease_score_distribution.append({
            "label": disease,
            "validAssessments": entry["count"],
            "dimensionMeans": [round(value / entry["count"], 2) for value in entry["sums"]],
            "totalMean": round(sum(entry["sums"]) / entry["count"], 2),
        })

    return {
        "sexDistribution": [{"label": key, "count": value} for key, value in sex_counts.most_common()],
        "ageDistribution": [{"label": key, "count": age_counts.get(key, 0)} for key in ["30岁及以下", "31–40岁", "41–50岁", "51–60岁", "61岁及以上"]],
        "regionDistribution": [{"label": key, "count": value} for key, value in region_counts.most_common()],
        "diseaseDistribution": [{"label": key, "count": value} for key, value in disease_counts.most_common()],
        "sexAgeDistribution": sex_age_distribution,
        "allergyDistribution": [{"label": key, "count": value} for key, value in allergy_counts.most_common()],
        "moduleCoverageByDisease": grouped_coverage("疾病"),
        "moduleCoverageByAge": grouped_coverage("年龄"),
        "healthPlanCoverage": ratio(len(health_plan_ids), patient_count),
        "trackingCoverage": ratio(len(tracking_ids), patient_count),
        "followupCoverage": ratio(len(followup_ids), patient_count),
        "symptomCoverage": ratio(len(symptom_ids), patient_count),
        "serviceExecution": {
            "healthPlans": len(data["healthPlans"]),
            "trackingPatients": len({normalize_patient_id(record) for record in tracking}),
            "medicationReminders": reminder_total,
            "temperatureMonitoring": temperature_total,
            "vitalMonitoring": vital_total,
            "totalPrompts": total_prompts,
            "estimatedResponses": round(weighted_response_numerator),
            "estimatedResponseRate": ratio(round(weighted_response_numerator), total_prompts),
            "responseRateBasis": "由患者级响应率按提醒总次数加权估算",
            "followupRecords": len(data["followups"]),
            "symptomRecords": len(data["symptomAssessments"]),
        },
        "followups": {"questions": followup_dimensions},
        "symptoms": {
            "questions": symptom_questions,
            "dimensionMeans": dimension_means,
            "validAssessments": len(total_scores),
            "scoreBands": [{"label": label, "count": score_bands[label]} for label in ["≤10分", "11–15分", "16–20分", "21–25分", ">25分"]],
            "diseaseDistribution": disease_score_distribution,
        },
        "medications": {
            "recordCount": len(medications),
            "uniqueDrugCount": len({clean(record.get("药品名称")) for record in medications if record.get("药品名称")}),
            "averageDrugsPerPatient": round(len(medications) / patient_count, 2) if patient_count else None,
            "drugDistribution": distribution(record.get("药品名称") for record in medications),
            "combinationCountDistribution": distribution(medication_patient_counts.values()),
            "productSpecificationDistribution": distribution(record.get("规格") for record in product_records),
            "productFrequencyDistribution": distribution(record.get("用药频率") for record in product_records),
            "productCourseDistribution": distribution(record.get("疗程天数") for record in product_records),
            "combinationModeDistribution": [{"label": label, "count": count} for label, count in combination_modes.most_common(12)],
        },
        "adverseEventRate": ratio(len(adverse_ids), patient_count) if adverse_events_provided else None,
        "adverseEvents": {
            "provided": adverse_events_provided,
            "recordCount": len(adverse_events) if adverse_events_provided else None,
            "patientCount": len(adverse_ids) if adverse_events_provided else None,
            "severityDistribution": distribution(record.get("不良反应严重程度分级") for record in adverse_events),
            "diseaseDistribution": distribution(record.get("疾病") for record in adverse_events),
            "outcomeDistribution": distribution(record.get("处理结果/转归") for record in adverse_events),
            "manualInterventionDistribution": distribution(record.get("是否触发人工干预") for record in adverse_events),
        },
        "riskDistribution": [
            {"label": "低风险", "count": low_risk},
            {"label": "中风险", "count": medium_risk},
            {"label": "高风险", "count": high_risk},
        ] if adverse_events_provided else [],
        "serviceGoals": [
            {"label": "健康管理方案覆盖", "actual": ratio(len(health_plan_ids), patient_count), "goal": "全量覆盖"},
            {"label": "智能随访覆盖", "actual": ratio(len(followup_ids), patient_count), "goal": "持续提升"},
            {"label": "症状自评覆盖", "actual": ratio(len(symptom_ids), patient_count), "goal": "持续提升"},
        ] + ([
            {"label": "不良反应监测", "actual": ratio(patient_count - len(adverse_ids), patient_count), "goal": "持续监测"},
        ] if adverse_events_provided else []),
    }


def _service_region_display(records):
    labels = []
    for record in records:
        label = clean(record.get("地区"))
        if label and label not in labels:
            labels.append(label)
    if not labels:
        return ""
    if len(labels) <= 3:
        return "、".join(labels)
    return f"共{len(labels)}个地区"


def build_insight(source_paths, product, start_text, end_text, client=None, provider=None):
    if len(source_paths) not in (4, 5, 6, 7):
        raise ValueError(f"月度洞察须提供4或5份Excel，兼容原6或7份资料，当前为{len(source_paths)}份")
    resolved = [str(Path(path).expanduser().resolve()) for path in source_paths]
    if len(set(resolved)) != len(resolved):
        raise ValueError("依据文件路径存在重复")
    start = date.fromisoformat(start_text)
    end = date.fromisoformat(end_text)
    if start > end:
        raise ValueError("服务周期开始日期不能晚于结束日期")

    data = {}
    diagnostics = {"roles": {}}
    for path in resolved:
        if not Path(path).is_file():
            raise ValueError(f"文件不存在：{path}")
        role, records, header_row = read_source(path)
        if role in data:
            raise ValueError(f"工作簿角色重复：{role}")
        included = filter_period(role, records, start, end)
        data[role] = included
        diagnostics["roles"][role] = {
            "path": path,
            "provided": True,
            "headerRow": header_row,
            "sourceRows": len(records),
            "includedRows": len(included),
        }

    monthly = "medicationReminders" in data
    required = {"patients", "medicationReminders", "followups", "symptomAssessments"} if monthly else REQUIRED_ROLES
    missing = sorted(required - set(data))
    if missing:
        raise ValueError(f"缺少工作簿角色：{', '.join(missing)}")

    patient_ids = [normalize_patient_id(record) for record in data["patients"]]
    if not patient_ids:
        raise ValueError("患者主表为空")
    if any(not patient_id for patient_id in patient_ids):
        raise ValueError("患者主表存在空userid")
    duplicates = [patient_id for patient_id, count in Counter(patient_ids).items() if count > 1]
    if duplicates:
        raise ValueError(f"患者主表userid重复：{duplicates[0]}")

    adverse_events_provided = "adverseEvents" in data
    if not adverse_events_provided:
        diagnostics["roles"]["adverseEvents"] = {
            "provided": False, "path": None, "headerRow": None,
            "sourceRows": None, "includedRows": None,
        }
    if monthly:
        from monthly_insight_metrics import build_monthly_metrics
        if set(data) - (required | {"adverseEvents"}):
            raise ValueError("月度五表模式不能混入旧版健康管理方案、跟踪提醒或逐药清单")
        known_ids = set(patient_ids)
        for role in data:
            if role != "patients":
                before = len(data[role])
                data[role] = [row for row in data[role] if normalize_patient_id(row) in known_ids]
                diagnostics["roles"][role]["unmatchedRows"] = before - len(data[role])
                diagnostics["roles"][role]["matchedRows"] = len(data[role])
        metrics = build_monthly_metrics(data, product, start, end)
    else:
        metrics = build_metrics(data, product)
    # Keep the record contract stable while availability distinguishes omission
    # from an explicitly supplied workbook with no in-period events.
    data.setdefault("adverseEvents", [])
    return {
        "schemaVersion": "1.0",
        "metadata": {
            "product": product,
            "inputMode": "monthly" if monthly else "legacy",
            "client": client or "",
            "provider": provider or "",
            "serviceRegion": _service_region_display(data["patients"]),
            "reportDate": f"{date.today().year}年{date.today().month}月",
            "period": {"start": start_text, "end": end_text},
            "patientCount": len(patient_ids),
            "regionCount": len(metrics["regionDistribution"]),
            "diseaseCount": len(metrics["diseaseDistribution"]),
            "contractGoal": "基于已提供的月度服务资料开展描述性评估，合同阈值未提供。" if monthly else "覆盖患者管理全流程，完成健康管理方案、用药提醒、AI智能随访、症状自评和不良反应监测，围绕产品安全性、有效性、依从性和患者体验进行服务评估。",
        },
        "sourceDiagnostics": diagnostics,
        "metrics": metrics,
        "records": data,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", action="append", required=True)
    parser.add_argument("--product", required=True)
    parser.add_argument("--start", required=True)
    parser.add_argument("--end", required=True)
    parser.add_argument("--client")
    parser.add_argument("--provider")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    result = build_insight(args.source, args.product, args.start, args.end, args.client, args.provider)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "patientCount": result["metadata"]["patientCount"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
