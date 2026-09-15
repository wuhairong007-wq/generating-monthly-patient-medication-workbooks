"""Nine-chapter report for monthly service exports, using source-supported measures."""
import json
from pathlib import Path
from docx import Document
from build_insight_report import (
    _configure_default_page, _clear_body, _clear_footer, _add_cover, _add_toc,
    _add_page_number, _set_update_fields, _add_heading, _add_para, _add_table,
    _add_caption, _add_chart,
)

CHAPTERS = ['一、报告概述', '二、患者基本特征分析', '三、AI用药提醒服务分析',
            '四、AI智能随访服务分析', '五、AI症状自评服务分析', '六、不良反应监测与分析',
            '七、患者风险评估', '八、服务质量与SLA达标分析', '九、总结']


def format_ratio(value):
    if value is None or value.get('value') is None:
        return '无法统计'
    return f"{value['numerator']:,}/{value['denominator']:,}（{value['value']:.1%}）"


def build_monthly_report(insight, chart_manifest_path, template_path, output_path):
    meta, m = insight['metadata'], insight['metrics']
    n, service, summary = meta['patientCount'], m['serviceExecution'], m['monthlySummary']
    doc = Document(template_path) if template_path else Document()
    if not template_path:
        _configure_default_page(doc)
    _clear_body(doc)
    for section in doc.sections:
        section.footer.is_linked_to_previous = False
        _clear_footer(section)
    _add_cover(doc, meta)
    _add_toc(doc, CHAPTERS)
    for section in doc.sections[:-1]:
        _clear_footer(section)
    _add_page_number(doc.sections[-1])
    _set_update_fields(doc)
    chart_map = {item['id']: item for item in json.loads(Path(chart_manifest_path).read_text(encoding='utf-8'))}
    chapter = figure_no = table_no = 0

    def heading(index):
        nonlocal chapter, figure_no, table_no
        chapter, figure_no, table_no = index, 0, 0
        _add_heading(doc, CHAPTERS[index-1], 1)

    def sub(text):
        _add_heading(doc, text, 2)

    def body(*paragraphs):
        for text in paragraphs:
            _add_para(doc, text)

    def table(headers, rows, caption):
        nonlocal table_no
        table_no += 1
        _add_table(doc, headers, rows or [['无记录'] + ['—'] * (len(headers)-1)])
        _add_caption(doc, '表', f'{chapter}-{table_no}', caption)

    def chart(key):
        nonlocal figure_no
        if key in chart_map:
            figure_no += 1
            item = chart_map[key]
            _add_chart(doc, Path(chart_manifest_path).parent/item['file'], item, f'{chapter}-{figure_no}')

    def distribution(key, caption, denominator=n):
        values = m.get(key, [])
        table(['类别', '人数', '占主表患者比例'],
              [[x['label'], x['count'], f"{x['count']/denominator:.1%}"] for x in values], caption)

    reminder_text = f"{service['medicationReminders']:,}次" if service['medicationReminders'] is not None else '无法统计'
    heading(1)
    sub('（一）执行摘要')
    body(f"本报告以月度患者清单中的{n:,}人为基础，分析{meta['period']['start']}至{meta['period']['end']}的服务登记。按患者标识关联随访、自评、用药方案及可选事件清单，提醒汇总须另外核验来源月份。",
         f"本周期可纳入的用药提醒总数为{reminder_text}，随访记录{service['followupRecords']:,}条，自评记录{service['symptomRecords']:,}条。记录量用于描述服务留痕，不直接等同于实际服药、症状改善或产品疗效。",
         '建议服务商先核对源表月份、患者归属与明细登记完整性，再按未覆盖患者和需人工复核的回答安排补访；委托方可据此跟踪资料补齐进度及后续服务动作。')
    sub('（二）关键指标概览')
    table(['指标', '本期统计', '口径'], [
        ['患者总数', f'{n:,}人', '月度主表唯一患者'],
        ['用药提醒总数', reminder_text, '完整同月主表汇总'],
        ['用药方案登记覆盖', format_ratio(m['reminderPlanCoverage']), '周期内确认且属于主表的患者'],
        ['智能随访覆盖', format_ratio(m['followupCoverage']), '周期内至少一次随访的患者'],
        ['症状自评覆盖', format_ratio(m['symptomCoverage']), '周期内至少一次自评的患者'],
    ], '关键服务指标与统计口径')
    body('覆盖率分母统一为本次患者主表唯一患者数；各模块独立去重，同一患者可同时出现在多个模块。不能将覆盖率相加，也不能把记录次数直接作为患者人数。')
    heading(2)
    features = [('（一）性别分布','sexDistribution','sex'), ('（二）年龄分布','ageDistribution','age'),
                ('（三）地区分布','regionDistribution',None), ('（四）疾病分型','diseaseDistribution','disease'),
                ('（五）过敏史','allergyDistribution','allergy')]
    for title, key, chart_key in features:
        sub(title)
        values = m[key]
        leader = max(values, key=lambda item: item['count']) if values else None
        body(f"本次主表按{title[3:]}汇总，以{n:,}名唯一患者为分母。" + (f"登记最多的类别为{leader['label']}，共{leader['count']:,}人。" if leader else '未取得可用分组。'),
             '该分布描述入组资料构成，不代表人群患病率或治疗效果；空字段保留为资料缺口，年龄未登记或无效者不应被解释为真实低龄患者。',
             '建议在下轮服务前复核主表对应字段，优先补齐缺失项；按疾病与年龄安排易理解的沟通材料，过敏史须由医护人员结合实际药物进一步核验。')
        distribution(key, title[3:])
        if chart_key: chart(chart_key)
    sub('（六）性别与年龄交叉')
    table(['性别','年龄段','人数'], [[x['sex'],x['age'],x['count']] for x in m['sexAgeDistribution']], '性别与年龄交叉分布')
    body('交叉分组以主表中的性别和有效年龄为依据，用于安排不同群体的服务资源。样本量较小的组不宜据此作效果比较；应先检查资料完整性，再决定是否调整联系时段。')
    chart('sex-age')
    heading(3)
    sub('（一）提醒次数与统计月份')
    body(f"源表标识月份为{summary['sourceMonth'] or '未识别'}，请求周期为{meta['period']['start']}至{meta['period']['end']}。本期用药提醒总数为{reminder_text}，仅在源表月份与请求的完整自然月一致时采用汇总字段。",
         '月度汇总没有逐次提醒日期，因此不能拆分到部分月份或迁移到另一个月。源月份不匹配、月份无法识别或提醒计数缺失时，将该指标保留为无法统计。',
         '建议服务商提供对应完整月份的主表或可按日期过滤的提醒明细，并核对总次数与患者级计数。当前资料不能支持提醒送达、点击或响应转化的推算。')
    chart('service')
    sub('（二）用药方案登记与周期')
    body(f"按用药方案确认时间筛选后，共有{service['reminderPlanRecords']:,}条登记，覆盖{format_ratio(m['reminderPlanCoverage'])}。用药周期来自源表明确字段，仅反映登记方案所载周期。",
         '方案确认记录并非逐次服药记录，也不证明患者已经开始或完成疗程。存在多次方案登记时应保留时间顺序，避免把先后调整误写为同一时间并用。',
         '建议复核未登记患者以及同一患者的多次确认记录，联系责任人员确认最新方案；用药调整需依据处方和医护意见，不由本报告生成新的剂量或频次。')
    table(['登记周期','方案记录数'], [[x['label'],x['count']] for x in m['medications']['planCycleDistribution']], '源表用药周期分布')
    sub('（三）登记药名与组合')
    body(f"联合用药字段共识别{m['medications']['uniqueDrugCount']:,}种登记药名，药名分布按患者与药名去重。当前产品名称精确匹配到{m['medications']['productPatientCount']:,}名登记患者；商品别名未自动合并。",
         '组合模式按每条方案登记计数，同一患者多次登记可能贡献多条记录。仅提取联合用药字段中的明确药名，不从用药方案叙述反推药品规格、实际频次或临床适宜性。',
         '建议对常见组合核对原始处方及变更时间，对名称不一致的条目先建立经确认的别名映射。高频登记组合可用于安排药师复核，不能直接证明联合用药的有效性或安全性。')
    table(['登记药名','对应患者数'], [[x['label'],x['count']] for x in m['medications']['drugDistribution']], '登记药名患者分布')
    chart('medications')
    table(['登记组合','方案记录数'], [[x['label'],x['count']] for x in m['medications']['combinationModeDistribution']], '常见方案组合')
    chart('combination-modes')
    heading(4)
    sub('（一）随访概况')
    body(f"本周期智能随访覆盖{format_ratio(m['followupCoverage'])}，共{service['followupRecords']:,}条记录。按随访时间与主表患者标识筛选后分析，重复随访作为独立答卷保留。",
         '各题分母为本题非空回答记录数，可能小于随访总记录数。A/B选项占比仅为选项统计，各题含义不同，不统一解释为正向回答率、依从率或满意度。',
         '建议先逐题查看实际选项含义，再对不适、未按方案执行或希望人工支持的回答安排复核；同一患者多次回答需要结合随访日期判断变化。')
    chart('followup')
    sub('（二）实际随访题目与回答')
    questions = m['followups']['questions']
    if not questions:
        body('本周期没有可供逐题分析的随访记录，无法给出回答分布或题目趋势。建议核对导出时间范围并补充对应月份资料；不得从其他月份复制回答或套用参考报告的题目。')
    for index, q in enumerate(questions,1):
        _add_heading(doc, q['question'], 3)
        body(f"本题非空回答共{q['positiveRate']['denominator']:,}条，A/B选项占比为{format_ratio(q['positiveRate'])}。具体选项及计数见下表，保持原始回答文字。",
             '解读应以本题提问和选项文字为准，同一字母在不同问题中不具有相同临床意义。缺失回答不计入本题分母，多次答卷也不等于新增患者。',
             '建议核查本题少见回答和需要人工解释的选项，按患者标识回溯原始随访记录；涉及病情变化时结合症状自评与事件登记安排复核，不直接认定产品因果关系。')
        table(['原始回答','记录数'], [[x['label'],x['count']] for x in q['distribution']], f'随访Q{index}回答分布')
    heading(5)
    sub('（一）自评概况与有效答卷')
    symptom = m['symptoms']
    body(f"本周期自评覆盖{format_ratio(m['symptomCoverage'])}，登记{service['symptomRecords']:,}条，其中六维度均可解析的有效答卷为{symptom['validAssessments']:,}条。",
         '评分按源表明确分值或A至E对应1至5分解析，六项完整才纳入均值与总分。分数只用于该问卷的描述性比较，需核查选项方向，不将其当作临床诊断或疗效证据。',
         '建议优先补齐不完整答卷，按原问题和选项核对高分回答，再与前后随访交叉复核。无有效答卷时停止评分比较，并补充对应周期的自评资料。')
    sub('（二）实际六维度评分')
    table(['源表问题','有效答卷均值'], [[q, '无法统计' if v is None else v] for q,v in zip(symptom['questions'],symptom['dimensionMeans'])], '自评各维度均值')
    chart('symptoms')
    sub('（三）按疾病分型评分分析')
    table(['疾病','有效答卷数','总分均值'], [[x['label'],x['validAssessments'],x['totalMean']] for x in symptom['diseaseDistribution']], '疾病分组自评总分')
    body('疾病分组的分母为该组完整答卷数，存在复测的患者可贡献多条答卷。不同疾病题意、样本规模与测量时点均可能影响比较，应核对原始资料后决定复测安排。')
    chart('disease-symptom')
    heading(6)
    sub('（一）事件登记概况')
    ae = m['adverseEvents']
    if ae['provided']:
        body(f"提供的不良反应清单中，本周期匹配主表患者的登记记录共{ae['recordCount']:,}条，涉及{ae['patientCount']:,}人。登记患者占比为{format_ratio(m['adverseEventRate'])}，分母为主表患者总数。",
             '上述数据反映清单登记情况，不等同于经过医学判定的药物不良反应发生率。零条登记不证明不存在事件；同一患者多条事件需分别保留时间和处理经过。',
             '建议责任人员核对发生日期、严重程度、处置与转归是否完整，并按患者标识与原始随访交叉确认；产品相关性须经个案评估，不由汇总比例推断。')
        table(['严重程度','事件记录数'], [[x['label'],x['count']] for x in ae['severityDistribution']], '事件严重程度登记')
        chart('adverse-events')
    else:
        body('本次未提供不良反应清单，无法评估事件登记规模、严重程度或转归。缺少清单属于资料可用性限制，不能解释为本周期已确认没有事件。',
             '建议补充带患者标识、发生日期、严重程度、处理措施和转归的登记表，再核对与随访、自评中不适描述的对应关系；补齐前保留统计空缺。')
    heading(7)
    sub('（一）服务复核优先级')
    if ae['provided']:
        table(['复核分组','患者数'], [[{'高风险':'优先复核','中风险':'其他事件患者','低风险':'未登记事件患者'}[x['label']],x['count']] for x in m['riskDistribution']], '基于事件登记的复核分组')
        body('人工干预或高度、重度、严重及危及生命事件患者列为优先复核，其余登记事件患者单列，其余主表患者仅表示没有匹配事件登记。本分组用于服务调度，不作为临床风险诊断。',
             '建议对优先复核组查阅处置记录和转归，对其他事件患者补齐随访；对未登记事件患者继续保留异常反馈入口，不能因清单中没有记录而认定安全。')
    else:
        body('未取得事件清单时，不输出基于事件的患者风险分布。可先依据实际随访问题和症状自评筛选需复核记录，但不得据此补造患者风险等级或事件发生情况。',
             '建议建立患者标识、异常反馈时间与处理结果之间的关联，补齐事件资料后再安排分层复核。医学紧急程度应由有资质的责任人员根据个案确定。')
    heading(8)
    sub('（一）模块覆盖与合同边界')
    table(['服务模块','患者覆盖','评价依据'], [[x['label'],format_ratio(x['actual']),x['goal']] for x in m['serviceGoals']], '现有资料支持的模块覆盖')
    body('本次仅展示用药方案登记、智能随访与症状自评的患者覆盖。未提供合同阈值、计划执行次数或送达响应日志，无法认定SLA达标，也不能以登记覆盖替代执行质量。',
         '建议委托方明确每项服务的考核分母、目标阈值与计划频次，服务商补充相应执行日志。后续评估应将合同目标与本期实际指标逐项对照，并记录未达项的原因。')
    sub('（二）疾病与年龄分组覆盖')
    for key, chart_key, caption in [('moduleCoverageByDisease','disease-coverage','疾病分组覆盖'),('moduleCoverageByAge','age-coverage','年龄分组覆盖')]:
        table(['分组','主表人数','方案登记','智能随访','症状自评'], [[x['label'],x['patientCount']]+[format_ratio(x['modules'][k]) for k in ['用药方案登记','智能随访','症状自评']] for x in m[key]],caption)
        chart(chart_key)
    body('分组覆盖率以各组主表唯一患者为分母，各模块按患者去重。低覆盖组应先检查导出资料是否齐全，再定位未登记人员及联系失败原因，避免把数据缺漏当作真实服务缺失。')
    sub('（三）源资料核验')
    roles = {'patients':'月度患者清单','medicationReminders':'用药提醒','followups':'患者随访','symptomAssessments':'症状自评','adverseEvents':'不良反应清单'}
    table(['资料角色','源记录数','周期内记录','匹配主表','未匹配主表'], [[roles[k],v.get('sourceRows') if v.get('provided') else '未提供',v.get('includedRows'),v.get('matchedRows',v.get('includedRows')),v.get('unmatchedRows',0) if v.get('provided') else '—'] for k,v in insight['sourceDiagnostics']['roles'].items()], '资料筛选与患者关联核验')
    body('日期筛选采用起止日期闭区间；主表作为患者分母，明细必须能匹配主表患者。未匹配记录不纳入统计，应回查标识格式和导出范围；源表月份与报告周期不一致时，月度提醒汇总不计入本期。')
    heading(9)
    body(f"本次报告围绕{n:,}名主表患者，形成用药方案登记、实际随访问答、症状自评与可用事件登记的描述性结果。用药提醒总数为{reminder_text}，各项结论受资料月份、患者关联和记录完整性约束。",
         '下一步优先补齐与服务周期一致的资料，复核未覆盖患者及异常回答，明确每项服务的责任人员和完成时限。对用药问题回到处方与医护评估，对登记缺口回到原始记录，不以汇总统计代替个体决策。',
         '下期可在相同患者范围、题目和统计口径下比较覆盖变化，先区分真实服务变化与导出资料差异。涉及改善、因果和达标的结论，须取得相应证据与评价标准后另行评估。')
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return str(output)
